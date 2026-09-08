import io
import re
import base64
from typing import BinaryIO, Any
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from markitdown import DocumentConverter, DocumentConverterResult
from markitdown.converters._pptx_converter import PptxConverter
from converter_engine.core.ocr_engine import extract_text_from_image

class PPTXParser(PptxConverter):
    """
    Custom PPTX Parser that subclasses MarkItDown's PptxConverter.
    It overrides the convert method to inject offline OCR using Tesseract
    for any embedded images or grouped shapes.
    """
    def convert(self, file_stream: BinaryIO, stream_info, **kwargs: Any) -> DocumentConverterResult:
        presentation = pptx.Presentation(file_stream)
        md_content = ""
        slide_num = 0
        
        for slide in presentation.slides:
            slide_num += 1
            md_content += f"\n\n<!-- Slide number: {slide_num} -->\n"
            title = slide.shapes.title if hasattr(slide.shapes, 'title') else None

            def _process_shape(shape):
                nonlocal md_content
                
                # Recursive group processing
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    sorted_shapes = sorted(
                        shape.shapes,
                        key=lambda x: (
                            float("-inf") if not x.top else x.top,
                            float("-inf") if not x.left else x.left,
                        ),
                    )
                    for subshape in sorted_shapes:
                        _process_shape(subshape)
                    return

                # Pictures and OCR
                if self._is_picture(shape):
                    image_blob, image_content_type, image_filename = self._get_image_info(shape)
                    
                    if image_blob:
                        ocr_text = extract_text_from_image(image_blob)
                        if ocr_text:
                            md_content += f"\n> **[Image Text Extraction]**\n"
                            for line in ocr_text.split("\n"):
                                if line.strip():
                                    md_content += f"> {line.strip()}\n"
                                    
                    # Fallback to standard alt text handling if keep_data_uris
                    alt_text = shape.name
                    try:
                        alt_text = shape._element._nvXxPr.cNvPr.attrib.get("descr", alt_text)
                    except Exception:
                        pass
                    
                    alt_text = re.sub(r"[\r\n\[\]]", " ", alt_text)
                    alt_text = re.sub(r"\s+", " ", alt_text).strip()
                    
                    if kwargs.get("keep_data_uris", False) and image_blob:
                        content_type = image_content_type or "image/png"
                        b64_string = base64.b64encode(image_blob).decode("utf-8")
                        md_content += f"\n![{alt_text}](data:{content_type};base64,{b64_string})\n"

                # Tables
                if self._is_table(shape):
                    md_content += self._convert_table_to_markdown(shape.table, **kwargs)

                # Charts
                if shape.has_chart:
                    md_content += self._convert_chart_to_markdown(shape.chart)

                # Text areas
                elif shape.has_text_frame:
                    if shape == title:
                        md_content += "# " + shape.text.lstrip() + "\n"
                    else:
                        md_content += shape.text + "\n"

            # Sort top-level shapes
            sorted_shapes = sorted(
                slide.shapes,
                key=lambda x: (
                    float("-inf") if not x.top else x.top,
                    float("-inf") if not x.left else x.left,
                ),
            )
            
            for shape in sorted_shapes:
                _process_shape(shape)

            md_content = md_content.strip()

            if slide.has_notes_slide:
                md_content += "\n\n### Notes:\n"
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame is not None:
                    md_content += notes_frame.text
                md_content = md_content.strip()

        return DocumentConverterResult(markdown=md_content.strip())
