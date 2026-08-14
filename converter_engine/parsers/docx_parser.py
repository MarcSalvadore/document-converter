"""Parser for DOCX documents using python-docx."""

import io
import os
from typing import BinaryIO, List, Optional, Union, Dict
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph, Run
from docx.oxml.ns import qn
from docx.oxml.shape import CT_Picture

from converter_engine.parsers import BaseParser, ParsedDocumentResult


class DOCXParser(BaseParser):
    """Parser for extracting Markdown from DOCX files."""

    def parse(self, source: Union[str, bytes, BinaryIO]) -> ParsedDocumentResult:
        """Parse DOCX document and return raw Markdown representation.

        Args:
            source: File path (str), raw bytes, or file-like binary stream.

        Returns:
            ParsedDocumentResult containing Markdown text and image bytes.
        """
        try:
            if isinstance(source, str):
                if not os.path.exists(source):
                    raise FileNotFoundError(f"DOCX file not found at: {source}")
                doc = docx.Document(source)
            elif isinstance(source, bytes):
                doc = docx.Document(io.BytesIO(source))
            else:
                doc = docx.Document(source)
        except FileNotFoundError:
            raise
        except Exception as e:
            raise ValueError(f"Failed to load DOCX document: {e}") from e


        md_blocks: List[str] = []
        extracted_images: Dict[str, bytes] = {}
        image_counter = [1] # Use list to allow mutation in nested methods

        # Iterate through body elements in document order
        for element in doc.element.body:
            if element.tag.endswith("p"):
                para = Paragraph(element, doc)
                parsed_para = self._parse_paragraph(para, doc, extracted_images, image_counter)
                if parsed_para.strip():
                    md_blocks.append(parsed_para)
            elif element.tag.endswith("tbl"):
                tbl = Table(element, doc)
                parsed_tbl = self._parse_table(tbl)
                if parsed_tbl.strip():
                    md_blocks.append(parsed_tbl)

        return ParsedDocumentResult(markdown="\n\n".join(md_blocks), images=extracted_images)

    def _parse_paragraph(self, para: Paragraph, doc: docx.Document, images: Dict[str, bytes], image_counter: List[int]) -> str:
        """Format paragraph into Markdown according to style and inline runs."""
        style_name = para.style.name if para.style else ""
        text = self._parse_runs_and_hyperlinks(para, doc, images, image_counter).strip()

        if not text:
            return ""

        # Heading style mapping
        if style_name.startswith("Heading"):
            level = self._extract_heading_level(style_name)
            prefix = "#" * level
            return f"{prefix} {text}"
        elif style_name == "Title":
            return f"# {text}"
        elif style_name == "Subtitle":
            return f"## {text}"

        # Bullet list mapping
        if "Bullet" in style_name or style_name.startswith("List Bullet"):
            indent_level = self._get_list_indent_level(para)
            indent = "  " * indent_level
            return f"{indent}- {text}"

        # Numbered list mapping
        if "Number" in style_name or style_name.startswith("List Number"):
            indent_level = self._get_list_indent_level(para)
            indent = "  " * indent_level
            return f"{indent}1. {text}"

        # Code block style
        if "Code" in style_name or "Source" in style_name:
            return f"```\n{text}\n```"

        return text

    def _extract_heading_level(self, style_name: str) -> int:
        """Extract heading level integer from style name like 'Heading 1'."""
        parts = style_name.split()
        if len(parts) > 1 and parts[-1].isdigit():
            level = int(parts[-1])
            return min(max(level, 1), 6)
        return 1

    def _get_list_indent_level(self, para: Paragraph) -> int:
        """Determine indent level for list items from style name or paragraph format."""
        style_name = para.style.name if para.style else ""
        parts = style_name.split()
        if len(parts) > 1 and parts[-1].isdigit():
            return max(0, int(parts[-1]) - 1)

        # Check numPr ilvl if available
        try:
            pPr = para._p.get_or_add_pPr()
            numPr = pPr.numPr
            if numPr is not None and numPr.ilvl is not None:
                return int(numPr.ilvl.val)
        except Exception:
            pass

        return 0

    def _parse_runs_and_hyperlinks(self, para: Paragraph, doc: docx.Document, images: Dict[str, bytes], image_counter: List[int]) -> str:
        """Extract paragraph text including inline formatting and hyperlinks, and extract images."""
        formatted_pieces: List[str] = []

        for child in para._p:
            if child.tag.endswith("r"):
                run = Run(child, para)
                formatted_pieces.append(self._format_run(run))
                
                # Check for images inside run
                for elem in child.iter():
                    if elem.tag.endswith("inline") or isinstance(elem, CT_Picture) or elem.tag.endswith("blip"):
                        embed_id = elem.attrib.get(qn('r:embed'))
                        if embed_id and embed_id in doc.part.rels:
                            rel = doc.part.rels[embed_id]
                            if "image" in rel.target_part.content_type:
                                img_part = rel.target_part
                                img_ext = img_part.partname.split('.')[-1]
                                idx = image_counter[0]
                                image_counter[0] += 1
                                filename = f"assets/docx_img_{idx}.{img_ext}"
                                images[filename] = img_part.blob
                                formatted_pieces.append(f"\n\n![Image {idx}]({filename})\n\n")

            elif child.tag.endswith("hyperlink"):
                # Extract text inside hyperlink node
                link_text_pieces = []
                r_id = child.attrib.get(qn("r:id"))
                url = None

                if r_id and r_id in doc.part.rels:
                    rel = doc.part.rels[r_id]
                    if rel.reltype == docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK:
                        url = rel.target_ref

                for sub_child in child:
                    if sub_child.tag.endswith("r"):
                        sub_run = Run(sub_child, para)
                        link_text_pieces.append(self._format_run(sub_run))

                link_text = "".join(link_text_pieces).strip()
                if url and link_text:
                    formatted_pieces.append(f"[{link_text}]({url})")
                elif link_text:
                    formatted_pieces.append(link_text)

        result = "".join(formatted_pieces)
        # Fallback to paragraph.text if run extraction yielded empty but paragraph text exists
        if not result and para.text:
            return para.text
        return result

    def _format_run(self, run: Run) -> str:
        """Apply bold, italic, code formatting to text run."""
        text = run.text
        if not text:
            return ""

        # Preserve spacing around formatted text
        l_space = " " if text.startswith(" ") and len(text) > 1 else ""
        r_space = " " if text.endswith(" ") and len(text) > 1 else ""
        core_text = text.strip()

        if not core_text:
            return text

        if run.font and run.font.name in ("Consolas", "Courier", "Courier New", "Monaco"):
            core_text = f"`{core_text}`"
        else:
            if run.bold and run.italic:
                core_text = f"***{core_text}***"
            elif run.bold:
                core_text = f"**{core_text}**"
            elif run.italic:
                core_text = f"*{core_text}*"

        return f"{l_space}{core_text}{r_space}"

    def _parse_table(self, table: Table) -> str:
        """Reconstruct table into Markdown pipe table format."""
        if not table.rows:
            return ""

        rows_data: List[List[str]] = []

        for row in table.rows:
            row_cells = []
            for cell in row.cells:
                # Replace internal newlines in cell with <br> and escape pipes
                cell_text = cell.text.replace("\n", "<br>").replace("|", "\\|").strip()
                row_cells.append(cell_text)
            rows_data.append(row_cells)

        if not rows_data:
            return ""

        # Build Markdown table
        header_row = rows_data[0]
        col_count = len(header_row)
        if col_count == 0:
            return ""

        separator_row = ["---"] * col_count

        md_table_lines = [
            "| " + " | ".join(header_row) + " |",
            "| " + " | ".join(separator_row) + " |",
        ]

        for row in rows_data[1:]:
            # Ensure row matches col_count length
            padded_row = row + [""] * (col_count - len(row))
            md_table_lines.append("| " + " | ".join(padded_row[:col_count]) + " |")

        return "\n".join(md_table_lines)
