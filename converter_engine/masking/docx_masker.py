# -*- coding: utf-8 -*-
import docx
from converter_engine.masking.patterns import find_matches, mask_text


def _process_paragraphs(paragraphs, company_pattern, opts, log, source_label):
    for i, para in enumerate(paragraphs):
        original = para.text
        if not original.strip():
            continue
        matches = find_matches(original, company_pattern, **opts)
        if not matches:
            continue
        new_text = mask_text(original, matches, mask_style="label")
        for m in matches:
            log.append({
                "lokasi": f"{source_label} paragraf #{i+1}",
                "label": m["label"],
                "teks_asli": m["text"],
            })
        # Tulis ulang isi paragraf ke run pertama, kosongkan run lainnya.
        # Ini mengorbankan sebagian format inline (bold/italic campuran di
        # satu kalimat) demi kepastian bahwa teks sensitif benar-benar hilang
        # dari XML, bukan cuma di run yang kebetulan cocok.
        if para.runs:
            para.runs[0].text = new_text
            for r in para.runs[1:]:
                r.text = ""
        else:
            para.add_run(new_text)


def redact_docx(input_path, output_path, company_names=None, **detect_opts):
    from converter_engine.masking.patterns import compile_company_list_pattern
    company_pattern = compile_company_list_pattern(company_names or [])
    doc = docx.Document(input_path)
    log = []

    _process_paragraphs(doc.paragraphs, company_pattern, detect_opts, log, "Body")

    for t_idx, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                _process_paragraphs(cell.paragraphs, company_pattern, detect_opts,
                                     log, f"Tabel #{t_idx+1}")

    for section in doc.sections:
        for header_footer, name in ((section.header, "Header"), (section.footer, "Footer")):
            _process_paragraphs(header_footer.paragraphs, company_pattern, detect_opts,
                                 log, name)

    doc.save(output_path)
    return log
