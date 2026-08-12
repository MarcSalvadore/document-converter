"""Shared Pytest fixtures and synthetic file generators."""

import io
import os
import pytest
import docx
import pptx
from pptx.util import Inches, Pt

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


@pytest.fixture
def sample_docx(tmp_path) -> str:
    """Generate a synthetic DOCX file with headings, formatted runs, lists, and tables."""
    file_path = str(tmp_path / "synthetic_sample.docx")
    doc = docx.Document()

    # Document Title and Headings
    doc.add_heading("Synthetic Document Title", level=0)
    doc.add_heading("Section One Header", level=1)
    doc.add_heading("Subsection Header", level=2)

    # Paragraph with inline formatting
    p = doc.add_paragraph("This paragraph contains ")
    r1 = p.add_run("bold text")
    r1.bold = True
    p.add_run(" and ")
    r2 = p.add_run("italic text")
    r2.italic = True
    p.add_run(" for testing parser extraction.")

    # Bullet List
    doc.add_paragraph("First bullet item", style="List Bullet")
    doc.add_paragraph("Second bullet item", style="List Bullet")

    # Numbered List
    doc.add_paragraph("First numbered item", style="List Number")
    doc.add_paragraph("Second numbered item", style="List Number")

    # Table
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Col 1"
    table.cell(0, 1).text = "Col 2"
    table.cell(0, 2).text = "Col 3"
    table.cell(1, 0).text = "Val A"
    table.cell(1, 1).text = "Val B"
    table.cell(1, 2).text = "Val C"

    doc.save(file_path)
    return file_path


@pytest.fixture
def sample_docx_bytes(sample_docx) -> bytes:
    """Return raw bytes of synthetic DOCX file."""
    with open(sample_docx, "rb") as f:
        return f.read()


@pytest.fixture
def sample_pptx(tmp_path) -> str:
    """Generate a synthetic PPTX file with two slides containing titles, bullets, and tables."""
    file_path = str(tmp_path / "synthetic_sample.pptx")
    prs = pptx.Presentation()

    # Slide 1: Title & Bullet list
    slide_layout_1 = prs.slide_layouts[1]  # Title and Content
    slide_1 = prs.slides.add_slide(slide_layout_1)
    slide_1.shapes.title.text = "Slide 1: Overview"

    tf_1 = slide_1.shapes.placeholders[1].text_frame
    tf_1.text = "Main bullet point"
    p_sub = tf_1.add_paragraph()
    p_sub.text = "Sub-bullet detail item"
    p_sub.level = 1

    # Slide 2: Title & Table
    blank_layout = prs.slide_layouts[6]
    slide_2 = prs.slides.add_slide(blank_layout)

    # Add title shape manually
    title_box = slide_2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    title_box.text_frame.text = "Slide 2: Summary Data"

    # Add 2x2 table shape
    table_shape = slide_2.shapes.add_table(rows=2, cols=2, left=Inches(0.5), top=Inches(2), width=Inches(6), height=Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "99%"

    prs.save(file_path)
    return file_path


@pytest.fixture
def sample_pptx_bytes(sample_pptx) -> bytes:
    """Return raw bytes of synthetic PPTX file."""
    with open(sample_pptx, "rb") as f:
        return f.read()


@pytest.fixture
def sample_pdf(tmp_path) -> str:
    """Generate a synthetic PDF file using ReportLab with headings, text, and tables."""
    file_path = str(tmp_path / "synthetic_sample.pdf")
    doc = SimpleDocTemplate(file_path, pagesize=letter)
    styles = getSampleStyleSheet()

    h1_style = ParagraphStyle("CustomH1", parent=styles["Heading1"], fontSize=20, leading=24)
    h2_style = ParagraphStyle("CustomH2", parent=styles["Heading2"], fontSize=15, leading=18)
    body_style = ParagraphStyle("CustomBody", parent=styles["Normal"], fontSize=10, leading=12)

    story = []
    story.append(Paragraph("PDF Report Title", h1_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph("Section One", h2_style))
    story.append(Spacer(1, 8))
    story.append(Paragraph("This is body paragraph content inside the generated PDF document.", body_style))
    story.append(Spacer(1, 12))

    # Table block
    table_data = [
        ["Header A", "Header B"],
        ["Cell 1", "Cell 2"],
    ]
    t = Table(table_data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
    ]))
    story.append(t)

    doc.build(story)
    return file_path


@pytest.fixture
def sample_pdf_bytes(sample_pdf) -> bytes:
    """Return raw bytes of synthetic PDF file."""
    with open(sample_pdf, "rb") as f:
        return f.read()


@pytest.fixture
def corrupted_file(tmp_path) -> str:
    """Generate a corrupted/invalid binary file with .docx extension."""
    file_path = str(tmp_path / "corrupted.docx")
    with open(file_path, "wb") as f:
        f.write(b"NOT_A_REAL_ZIP_CONTAINER_OR_DOCX")
    return file_path


@pytest.fixture
def unsupported_file(tmp_path) -> str:
    """Generate an unmapped .txt text file."""
    file_path = str(tmp_path / "notes.txt")
    with open(file_path, "w", encoding="utf-8") as f:
        f.write("Some standard plaintext content.")
    return file_path
