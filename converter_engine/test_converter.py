"""End-to-end test script to create sample documents and test DocumentConverter."""

import os
import sys
import docx
import pptx
from pptx.util import Inches, Pt

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# Ensure project root is in sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from converter_engine.core.router import DocumentRouter


def create_sample_docx(filename: str):
    doc = docx.Document()
    doc.add_heading("Sample Document Title", level=0)
    doc.add_heading("First Section Header", level=1)
    
    p1 = doc.add_paragraph("This is a standard paragraph with ")
    r1 = p1.add_run("bold text")
    r1.bold = True
    p1.add_run(" and ")
    r2 = p1.add_run("italic text")
    r2.italic = True
    p1.add_run(".")

    doc.add_heading("Subsection Header", level=2)
    doc.add_paragraph("Item 1", style="List Bullet")
    doc.add_paragraph("Item 2", style="List Bullet")
    doc.add_paragraph("Item 3", style="List Bullet")

    # Add table
    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "ID"
    hdr_cells[1].text = "Name"
    hdr_cells[2].text = "Role"

    row1 = table.rows[1].cells
    row1[0].text = "1"
    row1[1].text = "Alice Smith"
    row1[2].text = "Engineer"

    row2 = table.rows[2].cells
    row2[0].text = "2"
    row2[1].text = "Bob Jones"
    row2[2].text = "Architect"

    doc.save(filename)
    print(f"[TEST SETUP] Created sample DOCX: {filename}")


def create_sample_pptx(filename: str):
    prs = pptx.Presentation()
    slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(slide_layout)
    title1 = slide1.shapes.title
    title1.text = "Introduction Slide"
    subtitle1 = slide1.placeholders[1]
    subtitle1.text = "Welcome to the presentation"

    slide_layout2 = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(slide_layout2)
    title2 = slide2.shapes.title
    title2.text = "Key Agenda Items"
    tf2 = slide2.placeholders[1].text_frame
    tf2.text = "First agenda item"
    p2 = tf2.add_paragraph()
    p2.text = "Sub-bullet detail"
    p2.level = 1
    p3 = tf2.add_paragraph()
    p3.text = "Second agenda item"
    p3.level = 0

    prs.save(filename)
    print(f"[TEST SETUP] Created sample PPTX: {filename}")


def create_sample_pdf(filename: str):
    doc = SimpleDocTemplate(filename, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    h1_style = ParagraphStyle('H1', parent=styles['Heading1'], fontSize=22, leading=26)
    h2_style = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=16, leading=20)
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=10, leading=14)

    story.append(Paragraph("PDF Document Title", h1_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph("PDF Section Header", h2_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph("This is body text inside the PDF document.", body_style))
    story.append(Spacer(1, 12))

    data = [
        ["Project", "Status", "Owner"],
        ["Engine", "Complete", "Team A"],
        ["Parser", "In Progress", "Team B"]
    ]
    t = Table(data, colWidths=[150, 150, 150])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('GRID', (0,0), (-1,-1), 1, colors.black)
    ]))
    story.append(t)

    doc.build(story)
    print(f"[TEST SETUP] Created sample PDF: {filename}")


def run_tests():
    test_dir = os.path.join(os.path.dirname(__file__), "test_files")
    os.makedirs(test_dir, exist_ok=True)

    docx_path = os.path.join(test_dir, "sample.docx")
    pptx_path = os.path.join(test_dir, "sample.pptx")
    pdf_path = os.path.join(test_dir, "sample.pdf")

    create_sample_docx(docx_path)
    create_sample_pptx(pptx_path)
    create_sample_pdf(pdf_path)

    router = DocumentRouter()

    print("\n--- Testing DOCX Conversion ---")
    docx_md = router.convert(docx_path)
    print(docx_md)
    assert "# Sample Document Title" in docx_md
    assert "# First Section Header" in docx_md
    assert "| ID | Name | Role |" in docx_md
    print("DOCX Conversion PASSED!")

    print("\n--- Testing PPTX Conversion ---")
    pptx_md = router.convert(pptx_path)
    print(pptx_md)
    assert "Slide 1" in pptx_md
    assert "Key Agenda Items" in pptx_md
    print("PPTX Conversion PASSED!")

    print("\n--- Testing PDF Conversion ---")
    pdf_md = router.convert(pdf_path)
    print(pdf_md)
    assert "PDF Document Title" in pdf_md
    assert "PDF Section Header" in pdf_md
    assert "| Project | Status | Owner |" in pdf_md
    print("PDF Conversion PASSED!")

    print("\nALL CONVERTER ENGINE TESTS PASSED!")


if __name__ == "__main__":
    run_tests()
